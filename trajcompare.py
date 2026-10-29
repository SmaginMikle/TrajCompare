#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
trajcompare — автоматическое сравнение альтернативных траекторий скважин.

Режимы:
  pair   — сравнить одну пару файлов (эталон / альтернатива)
  batch  — пакетная обработка сотен/тысяч пар (параллельно, с журналом ошибок)

Примеры:
  python trajcompare.py pair Точное.xlsx Альтернативное.xlsx -o out
  python trajcompare.py batch --root wells/ -o out --workers 8
  python trajcompare.py batch --ref-dir ref/ --alt-dir alt/ -o out
  python trajcompare.py batch --manifest pairs.csv -o out
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import math
import os
import re
import sys
import time
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass, field, asdict
from pathlib import Path

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

log = logging.getLogger("trajcompare")

# ----------------------------------------------------------------------------
# Настройки (можно переопределить JSON-файлом через --config)
# ----------------------------------------------------------------------------
DEFAULTS = dict(
    step_md=5.0,            # шаг передискретизации по стволу, м
    wellhead_tol=1.0,       # допуск совпадения устьев, м
    target_tol=2.0,         # допуск попадания в цель, м
    sep_tol=10.0,           # допуск расхождения стволов для "Сходятся", м
    sep_warn=30.0,          # порог для "Условно сходятся", м
    landing_inc=80.0,       # зенит, с которого считаем участок горизонтальным, град
    kop_inc=0.5,            # зенит начала набора кривизны (KOP), град
    target_regex=r"(?<![A-Za-zА-Яа-я])[TТ]\s*\d+",  # T1/Т2/... (лат. или кирил.)
    plots=True,
    excel=True,
    dpi=130,
)

# Синонимы заголовков (нормализованные, нижний регистр). Порядок = приоритет.
COLS = {
    "md":       ["md", "глубина по стволу", "глубина по ств", "md (m)", "measured depth"],
    "inc":      ["incl", "inc", "зенитный угол", "зенит"],
    "azi_grid": ["азимут картограф", "azim grid", "grid az"],
    "azi_true": ["азимут истин", "azim true", "true az"],
    "azi":      ["azim", "азимут", "az"],
    "tvdss":    ["tvdss", "абс. отметка", "абс отметка", "абсолютная отметка", "z (m)"],
    "tvd":      ["tvd", "вертикаль"],
    "north":    ["x north", "north", "север", "x (m)"],
    "east":     ["y east", "east", "восток", "y (m)"],
    "comment":  ["коммент", "comment", "пласт", "target", "цель", "маркер"],
}
EXCLUDE = {  # чтобы "азимут" не схватил магнитный, а "север" — "Север/Восток" локальные
    "azi": ["магнит", "magn", "на точку"],
    "north": ["dx", "сю"],
    "east": ["dy", "вз"],
}


# ----------------------------------------------------------------------------
# Чтение файлов
# ----------------------------------------------------------------------------
def _norm(s) -> str:
    s = "" if s is None or (isinstance(s, float) and math.isnan(s)) else str(s)
    return re.sub(r"\s+", " ", s.replace("\n", " ")).strip().lower()


def _find_header_row(raw: pd.DataFrame, max_scan: int = 80) -> int:
    keys = [k for v in COLS.values() for k in v]
    best, best_row = 0, None
    for r in range(min(max_scan, len(raw))):
        cells = [_norm(c) for c in raw.iloc[r].tolist()]
        hits = sum(any(c.startswith(k) or k in c for k in keys) for c in cells if c)
        if hits > best:
            best, best_row = hits, r
    if best_row is None or best < 3:
        raise ValueError("не найдена строка заголовков (MD/зенит/азимут/координаты)")
    return best_row


def _pick(headers: list[str], data: pd.DataFrame, key: str):
    """Выбрать колонку по синонимам; при нескольких — с наибольшим числом значений."""
    cands = []
    for j, h in enumerate(headers):
        if not h:
            continue
        if any(x in h for x in EXCLUDE.get(key, [])):
            continue
        for pri, syn in enumerate(COLS[key]):
            if h.startswith(syn) or (len(syn) > 3 and syn in h):
                cands.append((pri, j))
                break
    if not cands:
        return None
    if key == "comment":
        return min(cands)[1]
    scored = []
    for pri, j in cands:
        v = pd.to_numeric(data.iloc[:, j], errors="coerce")
        scored.append((-v.notna().sum(), -float(np.nanmax(np.abs(v))) if v.notna().any() else 0, pri, j))
    return min(scored)[3]


@dataclass
class Trajectory:
    name: str
    path: str
    well: str
    md: np.ndarray
    inc: np.ndarray
    azi: np.ndarray
    n: np.ndarray       # абсолютная координата Север (X), м
    e: np.ndarray       # абсолютная координата Восток (Y), м
    z: np.ndarray       # TVDSS, м (вниз +)
    tvd: np.ndarray
    marks: list = field(default_factory=list)   # [(md, label)]
    azi_source: str = ""
    columns: dict = field(default_factory=dict)

    @property
    def xyz(self):
        return np.c_[self.n, self.e, self.z]


def min_curvature(md, inc, azi, n0=0.0, e0=0.0, z0=0.0):
    i, a = np.radians(inc), np.radians(azi)
    i1, i2, a1, a2 = i[:-1], i[1:], a[:-1], a[1:]
    cosb = np.cos(i2 - i1) - np.sin(i1) * np.sin(i2) * (1 - np.cos(a2 - a1))
    b = np.arccos(np.clip(cosb, -1, 1))
    rf = np.where(b > 1e-9, 2 / np.where(b > 1e-9, b, 1) * np.tan(b / 2), 1.0)
    dmd = np.diff(md)
    dn = dmd / 2 * (np.sin(i1) * np.cos(a1) + np.sin(i2) * np.cos(a2)) * rf
    de = dmd / 2 * (np.sin(i1) * np.sin(a1) + np.sin(i2) * np.sin(a2)) * rf
    dz = dmd / 2 * (np.cos(i1) + np.cos(i2)) * rf
    return (n0 + np.r_[0, np.cumsum(dn)], e0 + np.r_[0, np.cumsum(de)],
            z0 + np.r_[0, np.cumsum(dz)], np.degrees(b))


def read_trajectory(path: str | Path, name: str = "") -> Trajectory:
    path = Path(path)
    xl = pd.ExcelFile(path)
    sheet = xl.sheet_names[0]
    raw = pd.read_excel(xl, sheet_name=sheet, header=None)
    hr = _find_header_row(raw)
    headers = [_norm(c) for c in raw.iloc[hr].tolist()]
    data = raw.iloc[hr + 1:].reset_index(drop=True)

    cols = {k: _pick(headers, data, k) for k in COLS}
    if cols["azi_grid"] is not None:
        cols["azi"], azi_src = cols["azi_grid"], "картографический"
    elif cols["azi_true"] is not None and cols["azi"] in (None, cols["azi_true"]):
        cols["azi"], azi_src = cols["azi_true"], "истинный"
    else:
        azi_src = "как в файле"
    if cols["md"] is None or cols["inc"] is None or cols["azi"] is None:
        raise ValueError(f"нет обязательных колонок MD/зенит/азимут: {cols}")

    num = lambda k: pd.to_numeric(data.iloc[:, cols[k]], errors="coerce") if cols.get(k) is not None else None
    md = num("md")
    ok = md.notna()
    df = pd.DataFrame({"md": md, "inc": num("inc"), "azi": num("azi")})
    for k in ("tvdss", "tvd", "north", "east"):
        s = num(k)
        df[k] = s if s is not None else np.nan
    df["comment"] = data.iloc[:, cols["comment"]] if cols["comment"] is not None else ""
    df = df[ok].copy()
    df = df.drop_duplicates("md").sort_values("md").reset_index(drop=True)
    if len(df) < 3:
        raise ValueError("слишком мало точек траектории")

    md_, inc_, azi_ = df.md.values, df.inc.fillna(0).values, df.azi.ffill().fillna(0).values
    # Координаты: если в файле есть — берём их, иначе пересчёт минимальной кривизной
    if df.north.notna().all() and df.east.notna().all():
        n_, e_ = df.north.values, df.east.values
    else:
        n_, e_, _, _ = min_curvature(md_, inc_, azi_)
        n_ = n_ + (df.north.dropna().iloc[0] if df.north.notna().any() else 0)
        e_ = e_ + (df.east.dropna().iloc[0] if df.east.notna().any() else 0)
    if df.tvdss.notna().all():
        z_ = df.tvdss.values
        tvd_ = df.tvd.values if df.tvd.notna().all() else z_ - z_[0]
    elif df.tvd.notna().all():
        tvd_ = df.tvd.values
        z_ = tvd_.copy()
    else:
        _, _, tvd_, _ = min_curvature(md_, inc_, azi_)
        z_ = tvd_.copy()

    marks = []
    for m, c in zip(md_, df.comment.tolist()):
        c = _norm(c) and str(c).strip()
        if c and c.lower() != "nan":
            marks.append((float(m), c))

    well = sheet if sheet.lower() not in ("sheet1", "лист1") else path.stem
    return Trajectory(name or path.stem, str(path), str(well), md_, inc_, azi_,
                      n_, e_, z_, tvd_, marks, azi_src,
                      {k: (raw.iloc[hr, v] if v is not None else None) for k, v in cols.items()})


# ----------------------------------------------------------------------------
# Геометрия и метрики
# ----------------------------------------------------------------------------
def resample(t: Trajectory, step: float):
    grid = np.arange(t.md[0], t.md[-1], step)
    grid = np.r_[grid, t.md[-1]]
    return grid, np.c_[np.interp(grid, t.md, t.n), np.interp(grid, t.md, t.e),
                       np.interp(grid, t.md, t.z)]


def nearest_on_polyline(P, Q, mdQ, chunk=2000):
    """Для каждой точки P — ближайшая точка на ломаной Q (точно, до отрезка)."""
    A, B = Q[:-1], Q[1:]
    AB = B - A
    L2 = np.maximum((AB ** 2).sum(1), 1e-12)
    out_d = np.empty(len(P)); out_md = np.empty(len(P)); out_pt = np.empty_like(P)
    for s in range(0, len(P), chunk):
        p = P[s:s + chunk]
        AP = p[:, None, :] - A[None]
        t = np.clip((AP * AB[None]).sum(2) / L2[None], 0, 1)
        proj = A[None] + t[..., None] * AB[None]
        d = np.linalg.norm(p[:, None, :] - proj, axis=2)
        j = d.argmin(1)
        r = np.arange(len(p))
        out_d[s:s + chunk] = d[r, j]
        out_md[s:s + chunk] = mdQ[j] + t[r, j] * (mdQ[j + 1] - mdQ[j])
        out_pt[s:s + chunk] = proj[r, j]
    return out_d, out_md, out_pt


def dls(t: Trajectory):
    i, a = np.radians(t.inc), np.radians(t.azi)
    cosb = np.cos(i[1:] - i[:-1]) - np.sin(i[:-1]) * np.sin(i[1:]) * (1 - np.cos(a[1:] - a[:-1]))
    b = np.degrees(np.arccos(np.clip(cosb, -1, 1)))
    dmd = np.diff(t.md)
    return np.r_[0, np.where(dmd > 0, b / dmd * 10, 0)], b


def traj_params(t: Trajectory, cfg) -> dict:
    d10, beta = dls(t)
    hd = np.hypot(t.n - t.n[0], t.e - t.e[0])
    kop = t.md[np.argmax(t.inc > cfg["kop_inc"])] if (t.inc > cfg["kop_inc"]).any() else np.nan
    land_mask = t.inc >= cfg["landing_inc"]
    land = t.md[np.argmax(land_mask)] if land_mask.any() else np.nan
    # QC: пересчёт координат по MD/зениту/азимуту и подбор поправки азимута
    offs = np.arange(-6, 6.001, 0.02)
    best = (np.inf, 0.0)
    for o in offs:
        n_, e_, z_, _ = min_curvature(t.md, t.inc, t.azi + o, t.n[0], t.e[0], t.z[0])
        dev = np.hypot(n_ - t.n, e_ - t.e).max()
        if dev < best[0]:
            best = (dev, o)
    n0, e0, z0, _ = min_curvature(t.md, t.inc, t.azi, t.n[0], t.e[0], t.z[0])
    dev0 = np.sqrt((n0 - t.n) ** 2 + (e0 - t.e) ** 2 + (z0 - t.z) ** 2).max()
    return {
        "Длина по стволу, м": t.md[-1] - t.md[0],
        "Вертикальная глубина конца (TVDSS), м": t.z[-1],
        "Отход на забое, м": hd[-1],
        "Макс. зенит, град": float(np.nanmax(t.inc)),
        "Макс. интенсивность, град/10 м": float(np.nanmax(d10)),
        "Средняя интенсивность, град/10 м": float(np.nanmean(d10[1:])),
        "Суммарное искривление, град": float(beta.sum()),
        "Глубина зарезки (KOP), м": kop,
        f"MD входа в горизонт (зенит ≥ {cfg['landing_inc']:.0f}°), м": land,
        "Длина горизонтального участка, м": t.md[-1] - land if not np.isnan(land) else np.nan,
        "Число точек": len(t.md),
        "QC: отклонение пересчёта MC от координат файла, м": dev0,
        "QC: поправка азимута, совмещающая с координатами, град": best[1],
        "QC: остаток после поправки, м": best[0],
        "Тип азимута": t.azi_source,
    }


@dataclass
class PairResult:
    well: str
    ref: Trajectory
    alt: Trajectory
    summary: pd.DataFrame
    targets: pd.DataFrame
    profile: pd.DataFrame
    row: dict


def compare(ref: Trajectory, alt: Trajectory, cfg=DEFAULTS, well: str | None = None) -> PairResult:
    well = well or ref.well
    step = cfg["step_md"]
    gr, Pr = resample(ref, step)
    ga, Pa = resample(alt, step)

    # Расхождение: ближайшее расстояние (эталон → альтернатива и обратно)
    d_ra, md_on_a, pt_a = nearest_on_polyline(Pr, alt.xyz, alt.md)
    d_ar, _, _ = nearest_on_polyline(Pa, ref.xyz, ref.md)
    dv = pt_a[:, 2] - Pr[:, 2]
    dh = np.hypot(pt_a[:, 0] - Pr[:, 0], pt_a[:, 1] - Pr[:, 1])
    # Расхождение на одинаковой MD
    common = gr[gr <= alt.md[-1]]
    same = np.full(len(gr), np.nan)
    Pa_same = np.c_[np.interp(common, alt.md, alt.n), np.interp(common, alt.md, alt.e),
                    np.interp(common, alt.md, alt.z)]
    same[:len(common)] = np.linalg.norm(Pa_same - Pr[:len(common)], axis=1)

    profile = pd.DataFrame({
        "MD эталон, м": gr, "TVDSS эталон, м": Pr[:, 2],
        "MD альтернатива (ближ. точка), м": md_on_a,
        "Расхождение 3D, м": d_ra, "по горизонтали, м": dh, "по вертикали (alt−ref), м": dv,
        "ΔMD (alt−ref), м": md_on_a - gr, "Расхождение при равной MD, м": same,
    })

    # Реперные точки / цели: из комментариев обоих файлов + устье и забой
    rx = re.compile(cfg["target_regex"])
    pts = [("Устье", ref.md[0], ref, False), ("Забой эталона", ref.md[-1], ref, False)]
    for t in (ref, alt):
        for m, c in t.marks:
            pts.append((c, m, t, bool(rx.search(c))))
    if not any(p[3] for p in pts):           # нет подписанных целей → забой как цель
        pts[1] = ("Забой (цель по умолчанию)", ref.md[-1], ref, True)
    trows = []
    for label, m, src, is_t in pts:
        p = np.array([[np.interp(m, src.md, src.n), np.interp(m, src.md, src.e),
                       np.interp(m, src.md, src.z)]])
        other = alt if src is ref else ref
        d, mo, po = nearest_on_polyline(p, other.xyz, other.md)
        trows.append({
            "Точка": label, "Цель": "да" if is_t else "", "Источник": src.name,
            "MD в источнике, м": m, "TVDSS, м": p[0, 2],
            "Север (X), м": p[0, 0], "Восток (Y), м": p[0, 1],
            f"MD в {other.name}, м": mo[0], "ΔMD, м": mo[0] - m,
            "Промах 3D, м": d[0], "Промах по вертикали, м": po[0, 2] - p[0, 2],
            "Промах по горизонтали, м": float(np.hypot(*(po[0, :2] - p[0, :2]))),
        })
    targets = pd.DataFrame(trows)

    pr, pa = traj_params(ref, cfg), traj_params(alt, cfg)
    summary = pd.DataFrame({"Параметр": list(pr), ref.name: list(pr.values()), alt.name: list(pa.values())})
    summary["Разница (alt−ref)"] = [
        (b - a) if isinstance(a, (int, float, np.floating)) and isinstance(b, (int, float, np.floating)) else ""
        for a, b in zip(pr.values(), pa.values())]

    wh = float(np.linalg.norm(ref.xyz[0] - alt.xyz[0]))
    tg = targets[targets["Цель"] == "да"]
    tmiss = float(tg["Промах 3D, м"].max()) if len(tg) else np.nan
    haus = float(max(d_ra.max(), d_ar.max()))
    within = float((d_ra <= cfg["sep_tol"]).mean() * 100)
    imax = int(d_ra.argmax())

    reasons = []
    if wh > cfg["wellhead_tol"]:
        reasons.append(f"устья разнесены на {wh:.2f} м")
    if not np.isnan(tmiss) and tmiss > cfg["target_tol"]:
        reasons.append(f"промах по цели {tmiss:.2f} м")
    if haus > cfg["sep_tol"]:
        reasons.append(f"макс. расхождение {haus:.1f} м на MD≈{gr[imax]:.0f} м")
    if not reasons:
        verdict = "Сходятся"
    elif wh <= cfg["wellhead_tol"] and (np.isnan(tmiss) or tmiss <= cfg["target_tol"]) and haus <= cfg["sep_warn"]:
        verdict = "Условно сходятся"
    else:
        verdict = "Не сходятся"

    row = {
        "Скважина": well, "Эталон": Path(ref.path).name, "Альтернатива": Path(alt.path).name,
        "Вердикт": verdict, "Причины": "; ".join(reasons),
        "Расхождение устьев, м": wh, "Макс. промах по целям, м": tmiss,
        "Число целей": len(tg),
        "Макс. расхождение (Хаусдорф), м": haus,
        "MD макс. расхождения, м": gr[imax],
        "Среднее расхождение, м": float(d_ra.mean()),
        "P95 расхождения, м": float(np.percentile(d_ra, 95)),
        "RMS расхождения, м": float(np.sqrt((d_ra ** 2).mean())),
        f"Доля ствола в пределах {cfg['sep_tol']:.0f} м, %": within,
        "Длина эталона, м": pr["Длина по стволу, м"], "Длина альтернативы, м": pa["Длина по стволу, м"],
        "ΔДлины (alt−ref), м": pa["Длина по стволу, м"] - pr["Длина по стволу, м"],
        "Макс. DLS эталон, °/10м": pr["Макс. интенсивность, град/10 м"],
        "Макс. DLS альт., °/10м": pa["Макс. интенсивность, град/10 м"],
        "QC поправка азимута эталон, °": pr["QC: поправка азимута, совмещающая с координатами, град"],
        "QC поправка азимута альт., °": pa["QC: поправка азимута, совмещающая с координатами, град"],
    }
    return PairResult(well, ref, alt, summary, targets, profile, row)


# ----------------------------------------------------------------------------
# Графики
# ----------------------------------------------------------------------------
C_REF, C_ALT, C_SEP = "#1f5aa6", "#d9534f", "#6b6b6b"


def plot_pair(res: PairResult, out_png: Path, cfg=DEFAULTS):
    r, a = res.ref, res.alt
    n0, e0 = r.n[0], r.e[0]
    # Плоскость вертикального разреза — азимут на забой эталона
    az = math.atan2(r.e[-1] - e0, r.n[-1] - n0)
    sec = lambda t: (t.n - n0) * math.cos(az) + (t.e - e0) * math.sin(az)
    tg = res.targets[res.targets["Цель"] == "да"]

    fig = plt.figure(figsize=(20, 10.5))
    fig.suptitle(f"Скважина {res.well}: {r.name} vs {a.name} — {res.row['Вердикт']} "
                 f"(макс. расхождение {res.row['Макс. расхождение (Хаусдорф), м']:.1f} м)",
                 fontsize=14, weight="bold")
    gs = fig.add_gridspec(2, 4, hspace=0.32, wspace=0.28)

    def style(ax, title, xl, yl):
        ax.set_title(title, fontsize=11); ax.set_xlabel(xl); ax.set_ylabel(yl)
        ax.grid(alpha=.3)

    def mark_targets(ax, fx, fy):
        for _, row in tg.iterrows():
            x, y = fx(row), fy(row)
            ax.plot(x, y, marker="*", ms=13, color="#f0ad4e", mec="k", zorder=5)
            ax.annotate(str(row["Точка"]), (x, y), xytext=(5, 5), textcoords="offset points", fontsize=8)

    # 1. План
    ax = fig.add_subplot(gs[0, 0])
    ax.plot(r.e - e0, r.n - n0, color=C_REF, lw=2, label=r.name)
    ax.plot(a.e - e0, a.n - n0, color=C_ALT, lw=1.6, ls="--", label=a.name)
    ax.plot(0, 0, "ks", ms=6, label="Устье")
    mark_targets(ax, lambda q: q["Восток (Y), м"] - e0, lambda q: q["Север (X), м"] - n0)
    ax.set_aspect("equal", "datalim"); ax.legend(fontsize=8)
    style(ax, "План (горизонтальная проекция)", "Восток, м", "Север, м")

    # 2. Вертикальный разрез
    ax = fig.add_subplot(gs[0, 1])
    ax.plot(sec(r), r.z, color=C_REF, lw=2, label=r.name)
    ax.plot(sec(a), a.z, color=C_ALT, lw=1.6, ls="--", label=a.name)
    mark_targets(ax, lambda q: (q["Север (X), м"] - n0) * math.cos(az) + (q["Восток (Y), м"] - e0) * math.sin(az),
                 lambda q: q["TVDSS, м"])
    ax.invert_yaxis(); ax.legend(fontsize=8)
    style(ax, f"Вертикальный разрез (азимут {math.degrees(az) % 360:.1f}°)", "Отход вдоль разреза, м", "TVDSS, м")

    # 3. Проекция Север–TVDSS
    ax = fig.add_subplot(gs[0, 2])
    ax.plot(r.n - n0, r.z, color=C_REF, lw=2); ax.plot(a.n - n0, a.z, color=C_ALT, lw=1.6, ls="--")
    mark_targets(ax, lambda q: q["Север (X), м"] - n0, lambda q: q["TVDSS, м"])
    ax.invert_yaxis(); style(ax, "Проекция Север – глубина", "Север, м", "TVDSS, м")

    # 4. Проекция Восток–TVDSS
    ax = fig.add_subplot(gs[0, 3])
    ax.plot(r.e - e0, r.z, color=C_REF, lw=2); ax.plot(a.e - e0, a.z, color=C_ALT, lw=1.6, ls="--")
    mark_targets(ax, lambda q: q["Восток (Y), м"] - e0, lambda q: q["TVDSS, м"])
    ax.invert_yaxis(); style(ax, "Проекция Восток – глубина", "Восток, м", "TVDSS, м")

    # 5. 3D
    ax = fig.add_subplot(gs[1, 0], projection="3d")
    ax.plot(r.e - e0, r.n - n0, r.z, color=C_REF, lw=2)
    ax.plot(a.e - e0, a.n - n0, a.z, color=C_ALT, lw=1.6, ls="--")
    for _, q in tg.iterrows():
        ax.scatter(q["Восток (Y), м"] - e0, q["Север (X), м"] - n0, q["TVDSS, м"], marker="*", s=120,
                   color="#f0ad4e", edgecolor="k")
    ax.invert_zaxis(); ax.set_title("3D", fontsize=11); ax.zaxis.labelpad = 8
    ax.set_xlabel("Восток, м"); ax.set_ylabel("Север, м"); ax.set_zlabel("TVDSS, м")

    # 6. Расхождение по стволу
    ax = fig.add_subplot(gs[1, 1:3])
    p = res.profile
    ax.fill_between(p["MD эталон, м"], 0, p["Расхождение 3D, м"], color=C_SEP, alpha=.25)
    ax.plot(p["MD эталон, м"], p["Расхождение 3D, м"], color="k", lw=1.8, label="3D (ближайшее)")
    ax.plot(p["MD эталон, м"], p["по горизонтали, м"], color="#5cb85c", lw=1.2, label="по горизонтали")
    ax.plot(p["MD эталон, м"], p["по вертикали (alt−ref), м"], color="#9b59b6", lw=1.2, label="по вертикали (alt−ref)")
    ax.axhline(cfg["sep_tol"], color="#d9534f", ls=":", lw=1.2, label=f"допуск {cfg['sep_tol']:.0f} м")
    ax.axhline(0, color="k", lw=.5)
    for _, q in tg.iterrows():
        if q["Источник"] == r.name:
            ax.axvline(q["MD в источнике, м"], color="#f0ad4e", ls="--", lw=1)
            lbl = re.search(cfg["target_regex"], str(q["Точка"]))
            ax.text(q["MD в источнике, м"], ax.get_ylim()[1] * .95, lbl.group(0) if lbl else str(q["Точка"]),
                    fontsize=9, rotation=90, va="top", ha="right")
    ax.legend(fontsize=8, ncol=4, loc="upper left")
    style(ax, "Расхождение стволов по глубине эталона", "MD эталона, м", "м")

    # 7. Зенит и интенсивность
    ax = fig.add_subplot(gs[1, 3])
    ax.plot(r.md, r.inc, color=C_REF, lw=2, label=f"зенит {r.name}")
    ax.plot(a.md, a.inc, color=C_ALT, lw=1.4, ls="--", label=f"зенит {a.name}")
    ax.set_ylabel("Зенит, град")
    ax2 = ax.twinx()
    ax2.plot(r.md, dls(r)[0], color=C_REF, lw=.8, alpha=.5)
    ax2.plot(a.md, dls(a)[0], color=C_ALT, lw=.8, alpha=.5)
    ax2.set_ylabel("Интенсивность, град/10 м (тонкие)")
    ax.legend(fontsize=8, loc="lower right")
    style(ax, "Зенит и интенсивность по стволу", "MD, м", "Зенит, град")

    fig.savefig(out_png, dpi=cfg["dpi"], bbox_inches="tight")
    plt.close(fig)


# ----------------------------------------------------------------------------
# Excel-отчёт
# ----------------------------------------------------------------------------
VERDICT_FILL = {"Сходятся": "C6EFCE", "Условно сходятся": "FFEB9C", "Не сходятся": "FFC7CE", "Ошибка": "D9D9D9"}


def _format_ws(ws, widths=None, num_fmt="0.00"):
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5AA6")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"
    for col in ws.columns:
        L = col[0].column_letter
        w = (widths or {}).get(L) or min(max(len(str(col[0].value or "")) * 0.7, 10), 40)
        ws.column_dimensions[L].width = w
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = num_fmt


def write_pair_excel(res: PairResult, path: Path, png: Path | None):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import PatternFill
    ref, alt = res.ref, res.alt
    traj = lambda t: pd.DataFrame({"MD, м": t.md, "Зенит, град": t.inc, "Азимут, град": t.azi,
                                   "TVDSS, м": t.z, "Север (X), м": t.n, "Восток (Y), м": t.e,
                                   "Интенсивность, град/10м": dls(t)[0]})
    verdict = pd.DataFrame([{"Показатель": k, "Значение": v} for k, v in res.row.items()])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        verdict.to_excel(xw, "Итог", index=False)
        res.summary.to_excel(xw, "Параметры", index=False)
        res.targets.to_excel(xw, "Цели и реперы", index=False)
        res.profile.to_excel(xw, "Расхождение по стволу", index=False)
        traj(ref).to_excel(xw, f"Траектория {ref.name}"[:31], index=False)
        traj(alt).to_excel(xw, f"Траектория {alt.name}"[:31], index=False)
        wb = xw.book
        _format_ws(wb["Итог"], {"A": 45, "B": 55})
        _format_ws(wb["Параметры"], {"A": 58, "B": 20, "C": 20, "D": 20})
        _format_ws(wb["Цели и реперы"])
        for name in wb.sheetnames[3:]:
            _format_ws(wb[name])
        ws = wb["Итог"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "Вердикт":
                row[1].fill = PatternFill("solid", fgColor=VERDICT_FILL.get(row[1].value, "FFFFFF"))
        if png and png.exists():
            g = wb.create_sheet("Графики", 1)
            img = XLImage(str(png)); img.width, img.height = img.width * .55, img.height * .55
            g.add_image(img, "A1")


def write_batch_excel(rows: list[dict], errors: list[dict], path: Path, cfg):
    from openpyxl.styles import PatternFill
    df = pd.DataFrame(rows)
    stats = []
    if len(df):
        vc = df["Вердикт"].value_counts()
        for k in ("Сходятся", "Условно сходятся", "Не сходятся"):
            stats.append({"Показатель": k, "Значение": int(vc.get(k, 0))})
        stats += [
            {"Показатель": "Всего пар обработано", "Значение": len(df)},
            {"Показатель": "Ошибок чтения/расчёта", "Значение": len(errors)},
            {"Показатель": "Медиана макс. расхождения, м", "Значение": float(df["Макс. расхождение (Хаусдорф), м"].median())},
            {"Показатель": "P90 макс. расхождения, м", "Значение": float(df["Макс. расхождение (Хаусдорф), м"].quantile(.9))},
            {"Показатель": "Настройки", "Значение": json.dumps({k: v for k, v in cfg.items() if k not in ("plots", "excel")}, ensure_ascii=False)},
        ]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(stats).to_excel(xw, "Статистика", index=False)
        df.to_excel(xw, "Сравнение", index=False)
        pd.DataFrame(errors or [{"Скважина": "", "Ошибка": ""}]).to_excel(xw, "Ошибки", index=False)
        wb = xw.book
        _format_ws(wb["Статистика"], {"A": 40, "B": 60})
        _format_ws(wb["Сравнение"], {"A": 14, "B": 26, "C": 26, "D": 18, "E": 50})
        _format_ws(wb["Ошибки"], {"A": 20, "B": 40, "C": 40, "D": 80})
        ws = wb["Сравнение"]
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            row[3].fill = PatternFill("solid", fgColor=VERDICT_FILL.get(row[3].value, "FFFFFF"))


def plot_batch_overview(rows, path: Path, cfg):
    df = pd.DataFrame(rows)
    if df.empty:
        return
    fig, axs = plt.subplots(1, 3, figsize=(17, 4.8))
    h = df["Макс. расхождение (Хаусдорф), м"].clip(upper=df["Макс. расхождение (Хаусдорф), м"].quantile(.99))
    axs[0].hist(h, bins=40, color=C_REF, alpha=.8)
    axs[0].axvline(cfg["sep_tol"], color="r", ls=":"); axs[0].axvline(cfg["sep_warn"], color="orange", ls=":")
    axs[0].set_title("Макс. расхождение стволов"); axs[0].set_xlabel("м"); axs[0].set_ylabel("пар")
    axs[1].hist(df["Макс. промах по целям, м"].dropna(), bins=40, color=C_ALT, alpha=.8)
    axs[1].axvline(cfg["target_tol"], color="r", ls=":"); axs[1].set_title("Промах по целям"); axs[1].set_xlabel("м")
    vc = df["Вердикт"].value_counts().reindex(["Сходятся", "Условно сходятся", "Не сходятся"]).fillna(0)
    axs[2].bar(vc.index, vc.values, color=["#5cb85c", "#f0ad4e", "#d9534f"])
    axs[2].set_title("Вердикты")
    for ax in axs:
        ax.grid(alpha=.3)
    fig.tight_layout(); fig.savefig(path, dpi=120); plt.close(fig)


# ----------------------------------------------------------------------------
# Запуск: пара и пакет
# ----------------------------------------------------------------------------
def _safe(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", str(s)).strip("_") or "well"


def run_pair(ref_path, alt_path, out_dir, cfg=DEFAULTS, well=None, ref_name="Точное", alt_name="Альтернативное"):
    out_dir = Path(out_dir); out_dir.mkdir(parents=True, exist_ok=True)
    ref = read_trajectory(ref_path, ref_name)
    alt = read_trajectory(alt_path, alt_name)
    res = compare(ref, alt, cfg, well)
    stem = _safe(res.well)
    png = out_dir / f"{stem}_графики.png"
    if cfg["plots"]:
        plot_pair(res, png, cfg)
    if cfg["excel"]:
        write_pair_excel(res, out_dir / f"{stem}_сравнение.xlsx", png if cfg["plots"] else None)
    return res


def _worker(job):
    well, ref_p, alt_p, out_dir, cfg = job
    t0 = time.time()
    try:
        res = run_pair(ref_p, alt_p, Path(out_dir) / _safe(well), cfg, well)
        row = res.row; row["Время, с"] = round(time.time() - t0, 2)
        return True, row
    except Exception as ex:  # noqa: BLE001
        return False, {"Скважина": well, "Эталон": str(ref_p), "Альтернатива": str(alt_p),
                       "Ошибка": f"{type(ex).__name__}: {ex}", "trace": traceback.format_exc(limit=3)}


def _well_id(p: Path) -> str:
    try:
        s = pd.ExcelFile(p).sheet_names[0]
        if s.lower() not in ("sheet1", "лист1"):
            return s
    except Exception:  # noqa: BLE001
        pass
    return p.stem


def build_jobs(args, cfg):
    jobs = []
    rp, ap = re.compile(args.ref_pattern, re.I), re.compile(args.alt_pattern, re.I)
    if args.manifest:
        sep = ";" if ";" in Path(args.manifest).read_text(encoding="utf-8-sig").splitlines()[0] else ","
        with open(args.manifest, encoding="utf-8-sig") as f:
            for i, r in enumerate(csv.DictReader(f, delimiter=sep)):
                base = Path(args.manifest).parent
                rf, af = base / r["ref"], base / r["alt"]
                jobs.append((r.get("well") or _well_id(rf), rf, af, args.out, cfg))
    elif args.ref_dir and args.alt_dir:
        refs = {_well_id(p): p for p in Path(args.ref_dir).rglob("*.xls*") if not p.name.startswith("~$")}
        alts = {_well_id(p): p for p in Path(args.alt_dir).rglob("*.xls*") if not p.name.startswith("~$")}
        for w in sorted(set(refs) | set(alts)):
            if w in refs and w in alts:
                jobs.append((w, refs[w], alts[w], args.out, cfg))
            else:
                log.warning("Нет пары для скважины %s", w)
    elif args.root:
        groups: dict[Path, list[Path]] = {}
        for p in Path(args.root).rglob("*.xls*"):
            if not p.name.startswith("~$"):
                groups.setdefault(p.parent, []).append(p)
        for d, files in sorted(groups.items()):
            refs = [p for p in files if rp.search(p.name)]
            alts = [p for p in files if ap.search(p.name)]
            if len(refs) == 1 and len(alts) == 1:
                w = _well_id(refs[0])
                if Path(d) != Path(args.root):
                    w = f"{d.name}" if d.name != w else w
                jobs.append((w, refs[0], alts[0], args.out, cfg))
            else:
                # несколько пар в одной папке: сопоставляем по префиксу имени
                pref = lambda p, rx: rx.sub("", p.stem).strip(" _-")
                rmap = {pref(p, rp): p for p in refs}
                for p in alts:
                    k = pref(p, ap)
                    if k in rmap:
                        jobs.append((k or d.name, rmap[k], p, args.out, cfg))
                    else:
                        log.warning("Нет эталона для %s", p)
    else:
        raise SystemExit("Укажите --root, или --ref-dir и --alt-dir, или --manifest")
    return jobs


def run_batch(args, cfg, progress=None, cancel=None):
    out = Path(args.out); out.mkdir(parents=True, exist_ok=True)
    jobs = build_jobs(args, cfg)
    if args.skip_existing:
        jobs = [j for j in jobs if not (out / _safe(j[0]) / f"{_safe(j[0])}_сравнение.xlsx").exists()]
    log.info("Пар к обработке: %d, процессов: %d", len(jobs), args.workers)
    rows, errors = [], []
    t0 = time.time()
    if args.workers > 1:
        with ProcessPoolExecutor(args.workers) as ex:
            futs = [ex.submit(_worker, j) for j in jobs]
            for i, f in enumerate(as_completed(futs), 1):
                ok, r = f.result()
                (rows if ok else errors).append(r)
                if progress:
                    progress(i, len(jobs), r)
                if cancel and cancel():
                    for g in futs:
                        g.cancel()
                    break
                if i % 25 == 0 or i == len(jobs):
                    log.info("  %d/%d (%.0f с)", i, len(jobs), time.time() - t0)
    else:
        for i, j in enumerate(jobs, 1):
            if cancel and cancel():
                break
            ok, r = _worker(j)
            (rows if ok else errors).append(r)
            if progress:
                progress(i, len(jobs), r)
            log.info("  %d/%d %s: %s", i, len(jobs), j[0], r.get("Вердикт", r.get("Ошибка")))
    rows.sort(key=lambda r: str(r["Скважина"]))
    write_batch_excel(rows, errors, out / "СВОДНАЯ_сравнение_траекторий.xlsx", cfg)
    pd.DataFrame(rows).to_csv(out / "summary.csv", index=False, sep=";", encoding="utf-8-sig")
    plot_batch_overview(rows, out / "СВОДНАЯ_распределения.png", cfg)
    log.info("Готово: %d успешно, %d ошибок, %.1f с", len(rows), len(errors), time.time() - t0)
    return rows, errors


def main(argv=None):
    ap = argparse.ArgumentParser(description="Сравнение альтернативных траекторий скважин")
    sub = ap.add_subparsers(dest="cmd", required=True)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("-o", "--out", default="out")
    common.add_argument("--config", help="JSON с настройками (допуски, шаг, регулярка целей)")
    common.add_argument("--no-plots", action="store_true")
    for k in ("step_md", "target_tol", "sep_tol", "sep_warn", "wellhead_tol", "landing_inc"):
        common.add_argument("--" + k.replace("_", "-"), type=float)
    p1 = sub.add_parser("pair", parents=[common])
    p1.add_argument("ref"); p1.add_argument("alt"); p1.add_argument("--well")
    p2 = sub.add_parser("batch", parents=[common])
    p2.add_argument("--root", help="папка; в каждой подпапке пара файлов")
    p2.add_argument("--ref-dir"); p2.add_argument("--alt-dir")
    p2.add_argument("--manifest", help="CSV с колонками well;ref;alt")
    p2.add_argument("--ref-pattern", default=r"точн|exact|ref|факт|план")
    p2.add_argument("--alt-pattern", default=r"альтерн|alt")
    p2.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 2) - 1))
    p2.add_argument("--skip-existing", action="store_true", help="пропускать уже посчитанные")
    a = ap.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    cfg = dict(DEFAULTS)
    if a.config:
        cfg.update(json.loads(Path(a.config).read_text(encoding="utf-8")))
    for k in ("step_md", "target_tol", "sep_tol", "sep_warn", "wellhead_tol", "landing_inc"):
        if getattr(a, k, None) is not None:
            cfg[k] = getattr(a, k)
    if a.no_plots:
        cfg["plots"] = False

    if a.cmd == "pair":
        res = run_pair(a.ref, a.alt, a.out, cfg, a.well)
        for k, v in res.row.items():
            print(f"{k:45s} {v:.2f}" if isinstance(v, float) else f"{k:45s} {v}")
    else:
        run_batch(a, cfg)


if __name__ == "__main__":
    main()
