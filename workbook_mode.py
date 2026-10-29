# -*- coding: utf-8 -*-
"""
Режим «одна книга»:
  РН-ПД  — один Excel-файл, каждая скважина на своём листе (имя листа = номер скважины);
  РНТ — CSV (разделитель запятая): папка с файлом на скважину ИЛИ один CSV с колонкой «Скважина»;
  результат — один Excel-файл: лист «Сводная», затем по листу на скважину, в конце «Ошибки».
"""
from __future__ import annotations

import io
import logging
import re
import time
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import numpy as np
import pandas as pd

import trajcompare as tc

log = logging.getLogger("trajcompare")

# Кириллица/латиница одинакового начертания — чтобы «377Г-А» и «377Г-A» считались одной скважиной
_HOMO = str.maketrans("АВЕКМНОРСТХаеорсух", "ABEKMHOPCTXaeopcyx")


def well_key(s) -> str:
    return re.sub(r"[\s_\-\.]+", "", str(s)).translate(_HOMO).upper()


def _tokens(s) -> set[str]:
    parts = [p for p in re.split(r"[^0-9A-Za-zА-Яа-яЁё]+", str(s)) if p]
    toks = {well_key(p) for p in parts}
    for i in range(len(parts) - 1):          # «377Г_1» → тоже «377Г1»
        toks.add(well_key(parts[i] + parts[i + 1]))
    return toks


# ----------------------------------------------------------------------------
# Загрузка исходных данных
# ----------------------------------------------------------------------------
def load_reference_book(path) -> dict[str, pd.DataFrame]:
    log.info("Читаю книгу РН-ПД %s …", Path(path).name)
    t0 = time.time()
    book = pd.read_excel(path, sheet_name=None, header=None)
    book = {str(k).strip(): v for k, v in book.items() if v.dropna(how="all").shape[0] > 3}
    log.info("  листов со скважинами: %d (%.0f с)", len(book), time.time() - t0)
    return book


def load_alternatives(path) -> dict[str, tuple[str, pd.DataFrame]]:
    """Возвращает {имя скважины: (источник, сырая таблица)}."""
    path = Path(path)
    out: dict[str, tuple[str, pd.DataFrame]] = {}
    files = sorted(p for p in path.rglob("*.csv")) if path.is_dir() else [path]
    for f in files:
        try:
            raw = tc.read_csv_raw(f)
        except Exception as ex:  # noqa: BLE001
            log.warning("Не читается %s: %s", f.name, ex)
            out[f.stem] = (str(f), ex)          # ошибку покажем на листе «Ошибки»
            continue
        # Один CSV на много скважин — ищем колонку «Скважина»
        try:
            hr = tc._find_header_row(raw)
        except ValueError as ex:
            out[f.stem] = (str(f), ex)
            continue
        headers = [tc._norm(c) for c in raw.iloc[hr].tolist()]
        data = raw.iloc[hr + 1:]
        wc = tc._pick(headers, data, "well")
        if wc is not None and data.iloc[:, wc].nunique(dropna=True) > 1:
            for w, part in data.groupby(data.iloc[:, wc].astype(str).str.strip(), sort=False):
                if w and w.lower() != "nan":
                    out[w] = (f"{f.name} [{w}]", pd.concat([raw.iloc[[hr]], part]))
        else:
            w = str(data.iloc[:, wc].dropna().iloc[0]).strip() if wc is not None and data.iloc[:, wc].notna().any() else f.stem
            out[w if wc is not None else f.stem] = (f.name, raw)
    log.info("РНТ найдено: %d", len(out))
    return out


def match_wells(ref_names, alt_names):
    """Сопоставление листов РН-ПД и РНТ. Возвращает (пары, РН-ПД без пары, РНТ без пары)."""
    alt_by_key = {}
    for a in alt_names:
        alt_by_key.setdefault(well_key(a), []).append(a)
    used, pairs, no_alt = set(), [], []
    for r in ref_names:
        k = well_key(r)
        cand = [a for a in alt_by_key.get(k, []) if a not in used]
        if not cand:   # имя файла содержит номер скважины: «377Г_alt.csv», «РНТ 377Г.csv»
            cand = [a for a in alt_names if a not in used and k in _tokens(a)]
        if len(cand) >= 1:
            a = sorted(cand, key=len)[0]
            used.add(a)
            pairs.append((r, a))
        else:
            no_alt.append(r)
    no_ref = [a for a in alt_names if a not in used]
    return pairs, no_alt, no_ref


# ----------------------------------------------------------------------------
# Расчёт одной скважины (выполняется в отдельном процессе)
# ----------------------------------------------------------------------------
def _job(args):
    well, ref_raw, ref_src, alt_raw, alt_src, cfg = args
    t0 = time.time()
    try:
        ref = tc.parse_raw(ref_raw, "РН-ПД", ref_src, well)
        alt = tc.parse_raw(alt_raw, "РНТ", alt_src, well)
        res = tc.compare(ref, alt, cfg, well)
        png = None
        if cfg.get("plots", True):
            buf = io.BytesIO()
            tc.plot_pair(res, buf, cfg, dpi=cfg.get("dpi_embed", 80))
            try:   # 128 цветов: в 3 раза меньше, на графиках разницы не видно
                from PIL import Image
                im = Image.open(io.BytesIO(buf.getvalue())).convert("RGB").quantize(128)
                buf = io.BytesIO()
                im.save(buf, "PNG", optimize=True)
            except Exception:  # noqa: BLE001
                pass
            png = buf.getvalue()
        res.row["Время, с"] = round(time.time() - t0, 2)
        # Траектории в итоговую книгу не пишем — отдаём только то, что нужно
        light = dict(row=res.row, summary=res.summary, targets=res.targets,
                     profile=res.profile, png=png, ref_src=ref_src, alt_src=alt_src)
        return True, well, light
    except Exception as ex:  # noqa: BLE001
        return False, well, {"Скважина": well, "Проблема": "Ошибка расчёта",
                             "Подробности": f"{type(ex).__name__}: {ex}",
                             "РН-ПД": ref_src, "РНТ": alt_src,
                             "trace": traceback.format_exc(limit=3)}


# ----------------------------------------------------------------------------
# Итоговая книга
# ----------------------------------------------------------------------------
def _sheet_name(name, used):
    base = re.sub(r"[\[\]:*?/\\']", "_", str(name)).strip() or "Скважина"
    base = base[:31]
    n, i = base, 2
    while n.lower() in used:
        suf = f"_{i}"
        n = base[:31 - len(suf)] + suf
        i += 1
    used.add(n.lower())
    return n


def write_result_book(results: list[dict], errors: list[dict], out_path, cfg, sources: dict):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    BLUE = "1F5AA6"
    HFILL = PatternFill("solid", fgColor=BLUE)
    HFONT = Font(bold=True, color="FFFFFF")
    SEC = Font(bold=True, size=12, color=BLUE)
    thin = Side(style="thin", color="D0D0D0")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    VF = {k: PatternFill("solid", fgColor=v) for k, v in tc.VERDICT_FILL.items()}
    WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")

    def put_table(ws, r0, c0, df: pd.DataFrame, fmt="0.00"):
        for j, col in enumerate(df.columns):
            c = ws.cell(r0, c0 + j, str(col)); c.fill, c.font, c.alignment, c.border = HFILL, HFONT, WRAP, BOX
        ws.row_dimensions[r0].height = 46
        for i, row in enumerate(df.itertuples(index=False), 1):
            for j, v in enumerate(row):
                if isinstance(v, (float, np.floating)):
                    v = None if np.isnan(v) else float(v)
                elif isinstance(v, np.integer):
                    v = int(v)
                c = ws.cell(r0 + i, c0 + j, v)
                c.border = BOX
                if isinstance(v, float):
                    c.number_format = fmt
        return r0 + len(df) + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводная"
    used = {"сводная", "ошибки"}
    names = {}
    for res in results:
        names[res["row"]["Скважина"]] = _sheet_name(res["row"]["Скважина"], used)

    # ---------------- Сводная ----------------
    rows = [r["row"] for r in results]
    df = pd.DataFrame(rows)
    ws["A1"] = "Сравнение траекторий скважин: РН-ПД vs РНТ"
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    vc = df["Вердикт"].value_counts() if len(df) else pd.Series(dtype=int)
    stats = [
        ("РН-ПД (книга)", Path(sources.get("ref", "")).name),
        ("РНТ (CSV)", Path(sources.get("alt", "")).name),
        ("Сравнено скважин", len(df)),
        ("Сходятся", int(vc.get("Сходятся", 0))),
        ("Условно сходятся", int(vc.get("Условно сходятся", 0))),
        ("Не сходятся", int(vc.get("Не сходятся", 0))),
        ("Проблем (без пары / ошибки)", len(errors)),
        ("Допуски", f"цель {cfg['target_tol']:g} м, устье {cfg['wellhead_tol']:g} м, "
                    f"стволы {cfg['sep_tol']:g} / {cfg['sep_warn']:g} м"),
    ]
    for i, (k, v) in enumerate(stats, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        c.alignment = Alignment(horizontal="left")
        if k in VF:
            c.fill = VF[k]
    if errors:
        ws.cell(3 + len(stats) - 2, 3, "→ см. лист «Ошибки»").hyperlink = "#'Ошибки'!A1"
    r0 = 3 + len(stats) + 1
    cols_order = [c for c in ["Скважина", "Вердикт", "Причины"] if c in df.columns] + \
                 [c for c in df.columns if c not in ("Скважина", "Вердикт", "Причины", "РН-ПД", "РНТ")]
    if len(df):
        df = df[cols_order]
        last = put_table(ws, r0, 1, df)
        for i in range(r0 + 1, last):
            wcell = ws.cell(i, 1)
            sn = names.get(wcell.value)
            if sn:
                wcell.hyperlink = f"#'{sn}'!A1"
                wcell.font = Font(color="0563C1", underline="single", bold=True)
            v = ws.cell(i, 2)
            if v.value in VF:
                v.fill = VF[v.value]
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{last - 1}"
        ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.column_dimensions["A"].width = 34
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 48
    for j in range(4, len(cols_order) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15

    # ---------------- Листы скважин ----------------
    for k, res in enumerate(sorted(results, key=lambda r: str(r["row"]["Скважина"]))):
        row = res["row"]
        sh = wb.create_sheet(names[row["Скважина"]])
        v = row["Вердикт"]
        sh["A1"] = f"Скважина {row['Скважина']} — {v}"
        sh["A1"].font = Font(bold=True, size=14)
        sh["A1"].fill = VF.get(v, PatternFill())
        sh.merge_cells("A1:D1")
        sh["A2"] = "← к сводной"
        sh["A2"].hyperlink = "#'Сводная'!A1"
        sh["A2"].font = Font(color="0563C1", underline="single")
        sh["B2"] = f"РН-ПД: {res['ref_src']}"
        sh["B3"] = f"РНТ: {res['alt_src']}"

        r = 5
        sh.cell(r, 1, "Итог сравнения").font = SEC
        items = pd.DataFrame([{"Показатель": a, "Значение": b} for a, b in row.items()
                              if a not in ("Скважина", "РН-ПД", "РНТ")])
        r = put_table(sh, r + 1, 1, items)
        for i in range(7, r):
            if sh.cell(i, 1).value == "Вердикт":
                sh.cell(i, 2).fill = VF.get(v, PatternFill())
        r += 1
        sh.cell(r, 1, "Параметры траекторий").font = SEC
        r = put_table(sh, r + 1, 1, res["summary"])

        img_bottom = 5
        if res.get("png"):
            img = XLImage(io.BytesIO(res["png"]))
            scale = 1150 / img.width
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            sh.add_image(img, "F5")
            img_bottom = 5 + int(img.height / 20) + 2

        r = max(r + 1, img_bottom)
        sh.cell(r, 1, "Цели и реперные точки").font = SEC
        r = put_table(sh, r + 1, 1, res["targets"])
        r += 1
        step = cfg.get("profile_step", 25.0)
        p = res["profile"]
        md = p.iloc[:, 0].values
        keep = np.r_[True, np.diff(np.floor(md / step)) > 0]
        keep[-1] = True
        sh.cell(r, 1, f"Расхождение стволов по глубине (шаг {step:g} м)").font = SEC
        put_table(sh, r + 1, 1, p[keep].reset_index(drop=True))

        sh.column_dimensions["A"].width = 46
        for col in "BCD":
            sh.column_dimensions[col].width = 19
        for col in "EFGHIJKL":
            sh.column_dimensions[col].width = 14
        sh.sheet_view.zoomScale = 90
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        if (k + 1) % 100 == 0:
            log.info("  записано листов: %d", k + 1)

    # ---------------- Ошибки ----------------
    es = wb.create_sheet("Ошибки")
    edf = pd.DataFrame(errors or [{"Скважина": "", "Проблема": "нет", "Подробности": ""}])
    edf = edf[[c for c in ["Скважина", "Проблема", "Подробности", "РН-ПД", "РНТ"] if c in edf.columns]]
    put_table(es, 1, 1, edf)
    for col, w in zip("ABCDE", (22, 30, 80, 30, 30)):
        es.column_dimensions[col].width = w
    es["G1"] = "← к сводной"
    es["G1"].hyperlink = "#'Сводная'!A1"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# ----------------------------------------------------------------------------
# Запуск
# ----------------------------------------------------------------------------
def run_workbook(ref_xlsx, alt_csv, out_xlsx, cfg=None, workers=1, progress=None, cancel=None):
    cfg = dict(tc.DEFAULTS, **(cfg or {}))
    t0 = time.time()
    book = load_reference_book(ref_xlsx)
    alts = load_alternatives(alt_csv)
    pairs, no_alt, no_ref = match_wells(list(book), list(alts))
    log.info("Пар: %d; РН-ПД без РНТ: %d; РНТ без РН-ПД: %d",
             len(pairs), len(no_alt), len(no_ref))

    errors = [{"Скважина": w, "Проблема": "Нет CSV РНТ", "Подробности":
               "лист есть в книге РН-ПД, подходящий CSV не найден", "РН-ПД": f"лист «{w}»"} for w in no_alt]
    jobs = []
    for r, a in pairs:
        src, raw = alts[a]
        if isinstance(raw, Exception):
            errors.append({"Скважина": r, "Проблема": "CSV не читается", "Подробности": str(raw), "РНТ": src})
            continue
        jobs.append((r, book[r], f"{Path(ref_xlsx).name} / лист «{r}»", raw, src, cfg))
    for a in no_ref:
        src, raw = alts[a]
        errors.append({"Скважина": a, "Проблема": "Нет листа в книге РН-ПД" if not isinstance(raw, Exception)
                       else "CSV не читается", "Подробности": "" if not isinstance(raw, Exception) else str(raw),
                       "РНТ": src})

    results = []
    n = len(jobs)

    def handle(i, ok, well, data):
        if ok:
            results.append(data)
            r = data["row"]
        else:
            errors.append(data)
            r = {"Скважина": well, "Ошибка": data["Подробности"]}
        if progress:
            progress(i, n, r)

    if workers > 1 and n > 1:
        with ProcessPoolExecutor(workers) as ex:
            futs = [ex.submit(_job, j) for j in jobs]
            for i, f in enumerate(as_completed(futs), 1):
                handle(i, *f.result())
                if cancel and cancel():
                    for g in futs:
                        g.cancel()
                    break
                if i % 25 == 0 or i == n:
                    log.info("  посчитано %d/%d (%.0f с)", i, n, time.time() - t0)
    else:
        for i, j in enumerate(jobs, 1):
            if cancel and cancel():
                break
            handle(i, *_job(j))

    log.info("Записываю итоговую книгу …")
    out = write_result_book(results, errors, out_xlsx, cfg, {"ref": ref_xlsx, "alt": alt_csv})
    log.info("Готово: %s — скважин %d, проблем %d, %.0f с", out.name, len(results), len(errors), time.time() - t0)
    return results, errors, out
